# ============================================================
# Script: gerar_relatorio_ad.py
# Descricao: Le o CSV gerado pelo HealthCheckAD e gera um
#            relatorio DOCX profissional com formatacao,
#            tabelas coloridas, graficos e sumario executivo.
# Requer: pip install python-docx matplotlib
# Uso: python gerar_relatorio_ad.py <arquivo.csv>
# ============================================================

import csv
import sys
import os
from datetime import datetime
from collections import Counter
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Inches, Pt, Cm, RGBColor, Emu
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.section import WD_ORIENT
    from docx.oxml.ns import qn, nsdecls
    from docx.oxml import parse_xml
except ImportError:
    print("ERRO: Instale python-docx -> pip install python-docx")
    sys.exit(1)

try:
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    HAS_MATPLOTLIB = True
except ImportError:
    HAS_MATPLOTLIB = False
    print("AVISO: matplotlib nao encontrado. Graficos serao ignorados.")
    print("       Instale com: pip install matplotlib")


# ============================================================
# CORES E CONSTANTES
# ============================================================
CORES = {
    'OK':       RGBColor(0x28, 0xA7, 0x45),
    'AVISO':    RGBColor(0xFF, 0xC1, 0x07),
    'ERRO':     RGBColor(0xDC, 0x35, 0x45),
    'INFO':     RGBColor(0x17, 0xA2, 0xB8),
}

CORES_HEX = {
    'OK':    '#28a745',
    'AVISO': '#ffc107',
    'ERRO':  '#dc3545',
    'INFO':  '#17a2b8',
}

CORES_BG = {
    'OK':    'D4EDDA',
    'AVISO': 'FFF3CD',
    'ERRO':  'F8D7DA',
    'INFO':  'D1ECF1',
}

COR_HEADER_TABELA = '0078D4'
COR_HEADER_TEXTO = 'FFFFFF'
COR_LINHA_ALTERNADA = 'F2F7FB'


# ============================================================
# FUNCOES AUXILIARES
# ============================================================
def ler_csv(caminho: str) -> list[dict]:
    """Le o CSV com deteccao automatica de delimitador."""
    with open(caminho, 'r', encoding='utf-8-sig') as f:
        amostra = f.read(2048)
        f.seek(0)

        # Detectar delimitador
        if amostra.count(';') > amostra.count(','):
            delimitador = ';'
        else:
            delimitador = ','

        reader = csv.DictReader(f, delimiter=delimitador)

        # Validar colunas esperadas
        colunas_esperadas = {'Categoria', 'Teste', 'Status', 'Detalhe'}
        colunas_encontradas = set(reader.fieldnames or [])

        if not colunas_esperadas.issubset(colunas_encontradas):
            faltando = colunas_esperadas - colunas_encontradas
            print(f"ERRO: Colunas ausentes no CSV: {faltando}")
            print(f"       Colunas encontradas: {colunas_encontradas}")
            sys.exit(1)

        dados = [row for row in reader]

    if not dados:
        print("ERRO: CSV vazio.")
        sys.exit(1)

    return dados


def contar_status(dados: list[dict]) -> dict:
    """Conta ocorrencias de cada status."""
    return Counter(row.get('Status', '').strip().upper() for row in dados)


def agrupar_por_categoria(dados: list[dict]) -> dict:
    """Agrupa linhas por Categoria mantendo a ordem original."""
    grupos = {}
    for row in dados:
        cat = row.get('Categoria', 'Outros').strip()
        if cat not in grupos:
            grupos[cat] = []
        grupos[cat].append(row)
    return grupos


def definir_cor_celula(cell, cor_hex: str):
    """Aplica cor de fundo a uma celula da tabela."""
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{cor_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def definir_bordas_tabela(table):
    """Aplica bordas finas a toda a tabela."""
    tbl = table._tbl
    tbl_pr = tbl.tblPr if tbl.tblPr is not None else parse_xml(f'<w:tblPr {nsdecls("w")}/>')

    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        '  <w:top w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        '  <w:left w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        '  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        '  <w:right w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        '  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        '  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        '</w:tblBorders>'
    )
    tbl_pr.append(borders)


def formatar_header_tabela(row, num_colunas: int):
    """Formata a linha de cabecalho da tabela."""
    for i in range(num_colunas):
        cell = row.cells[i]
        definir_cor_celula(cell, COR_HEADER_TABELA)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                run.font.bold = True
                run.font.size = Pt(9)


def adicionar_paragrafo(doc, texto: str, estilo=None, negrito=False,
                         tamanho=None, cor=None, alinhamento=None,
                         espacamento_antes=None, espacamento_depois=None):
    """Adiciona paragrafo com formatacao."""
    p = doc.add_paragraph(style=estilo)
    run = p.add_run(texto)

    if negrito:
        run.bold = True
    if tamanho:
        run.font.size = Pt(tamanho)
    if cor:
        run.font.color.rgb = cor
    if alinhamento:
        p.alignment = alinhamento
    if espacamento_antes is not None:
        p.paragraph_format.space_before = Pt(espacamento_antes)
    if espacamento_depois is not None:
        p.paragraph_format.space_after = Pt(espacamento_depois)

    return p


def gerar_grafico_pizza(contagem: dict, caminho_img: str):
    """Gera grafico de pizza com os status."""
    if not HAS_MATPLOTLIB:
        return False

    labels = []
    sizes = []
    colors = []

    ordem = ['OK', 'AVISO', 'ERRO', 'INFO']
    for status in ordem:
        qtd = contagem.get(status, 0)
        if qtd > 0:
            labels.append(f"{status} ({qtd})")
            sizes.append(qtd)
            colors.append(CORES_HEX.get(status, '#6c757d'))

    if not sizes:
        return False

    fig, ax = plt.subplots(figsize=(5, 3.5))
    wedges, texts, autotexts = ax.pie(
        sizes,
        labels=labels,
        colors=colors,
        autopct='%1.1f%%',
        startangle=90,
        pctdistance=0.75,
        textprops={'fontsize': 9}
    )

    for at in autotexts:
        at.set_fontsize(8)
        at.set_color('white')
        at.set_fontweight('bold')

    ax.set_title('Distribuicao dos Resultados', fontsize=12, fontweight='bold', pad=15)
    plt.tight_layout()
    plt.savefig(caminho_img, dpi=150, bbox_inches='tight', facecolor='white')
    plt.close()
    return True


def gerar_grafico_barras_categoria(dados: list[dict], caminho_img: str):
    """Gera grafico de barras com erros/avisos por categoria."""
    if not HAS_MATPLOTLIB:
        return False

    categorias = {}
    for row in dados:
        cat = row.get('Categoria', 'Outros').strip()
        status = row.get('Status', '').strip().upper()
        if cat not in categorias:
            categorias[cat] = {'OK': 0, 'AVISO': 0, 'ERRO': 0, 'INFO': 0}
        if status in categorias[cat]:
            categorias[cat][status] += 1

    # Filtrar categorias que tem erros ou avisos
    cats_relevantes = {k: v for k, v in categorias.items()
                       if v['ERRO'] > 0 or v['AVISO'] > 0}

    if not cats_relevantes:
        # Mostrar top 10 por total
        cats_relevantes = dict(sorted(categorias.items(),
                                       key=lambda x: sum(x[1].values()),
                                       reverse=True)[:10])

    if not cats_relevantes:
        return False

    nomes = list(cats_relevantes.keys())
    erros = [cats_relevantes[c]['ERRO'] for c in nomes]
    avisos = [cats_relevantes[c]['AVISO'] for c in nomes]
    oks = [cats_relevantes[c]['OK'] for c in nomes]

    # Truncar nomes longos
    nomes_curtos = [n[:25] + '...' if len(n) > 28 else n for n in nomes]

    fig, ax = plt.subplots(figsize=(7, max(3, len(nomes) * 0.4)))

    y_pos = range(len(nomes))
    bars_err = ax.barh(y_pos, erros, color='#dc3545', label='ERRO', height=0.25)
    bars_avs = ax.barh([y + 0.25 for y in y_pos], avisos, color='#ffc107', label='AVISO', height=0.25)
    bars_ok = ax.barh([y + 0.5 for y in y_pos], oks, color='#28a745', label='OK', height=0.25)

    ax.set_yticks([y + 0.25 for y in y_pos])
    ax.set_yticklabels(nomes_curtos, fontsize=8)
    ax.set_xlabel('Quantidade', fontsize=9)
    ax.set_title('Resultados por Categoria', fontsize=12, fontweight='bold')
    ax.legend(fontsize=8, loc='lower right')
    ax.invert_yaxis()

    plt.tight_layout()
    plt.savefig(caminho_img, dpi=150, bbox_inches='tight', facecolor='white')
    plt.close()
    return True


# ============================================================
# GERACAO DO DOCUMENTO
# ============================================================
def gerar_docx(dados: list[dict], caminho_saida: str):
    """Gera o documento DOCX completo."""

    doc = Document()

    # --- Configurar estilos padrao ---
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(10)
    style.paragraph_format.space_after = Pt(4)

    # --- Configurar margens ---
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)

    contagem = contar_status(dados)
    grupos = agrupar_por_categoria(dados)
    total = len(dados)

    total_ok = contagem.get('OK', 0)
    total_avisos = contagem.get('AVISO', 0)
    total_erros = contagem.get('ERRO', 0)
    total_info = contagem.get('INFO', 0)

    if total_erros > 0:
        saude_geral = "AD COM PROBLEMAS"
        cor_saude = CORES['ERRO']
    elif total_avisos > 0:
        saude_geral = "AD OK COM AVISOS"
        cor_saude = CORES['AVISO']
    else:
        saude_geral = "AD SAUDAVEL"
        cor_saude = CORES['OK']

    # ========================================
    # CAPA
    # ========================================
    for _ in range(4):
        doc.add_paragraph()

    adicionar_paragrafo(doc, "RELATORIO DE SAUDE",
                         negrito=True, tamanho=28,
                         cor=RGBColor(0x00, 0x78, 0xD4),
                         alinhamento=WD_ALIGN_PARAGRAPH.CENTER)

    adicionar_paragrafo(doc, "Active Directory",
                         negrito=True, tamanho=22,
                         cor=RGBColor(0x1A, 0x1A, 0x2E),
                         alinhamento=WD_ALIGN_PARAGRAPH.CENTER)

    doc.add_paragraph()

    adicionar_paragrafo(doc, f"Status: {saude_geral}",
                         negrito=True, tamanho=16,
                         cor=cor_saude,
                         alinhamento=WD_ALIGN_PARAGRAPH.CENTER)

    doc.add_paragraph()
    doc.add_paragraph()

    adicionar_paragrafo(doc, f"Data de Geracao: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}",
                         tamanho=11,
                         alinhamento=WD_ALIGN_PARAGRAPH.CENTER)

    adicionar_paragrafo(doc, f"Total de Verificacoes: {total}",
                         tamanho=11,
                         alinhamento=WD_ALIGN_PARAGRAPH.CENTER)

    adicionar_paragrafo(doc, "Documento Confidencial",
                         negrito=True, tamanho=10,
                         cor=RGBColor(0xDC, 0x35, 0x45),
                         alinhamento=WD_ALIGN_PARAGRAPH.CENTER,
                         espacamento_antes=40)

    doc.add_page_break()

    # ========================================
    # SUMARIO (indice)
    # ========================================
    doc.add_heading('Sumario', level=1)

    itens_sumario = [
        "1. Sumario Executivo",
        "2. Estatisticas Gerais",
        "3. Itens Criticos (Erros)",
        "4. Itens de Atencao (Avisos)",
        "5. Resultados por Categoria",
        "6. Tabela Completa de Resultados",
        "7. Recomendacoes",
    ]
    for item in itens_sumario:
        adicionar_paragrafo(doc, item, tamanho=11, espacamento_antes=2, espacamento_depois=2)

    doc.add_page_break()

    # ========================================
    # 1. SUMARIO EXECUTIVO
    # ========================================
    doc.add_heading('1. Sumario Executivo', level=1)

    p = doc.add_paragraph()
    p.add_run('Este relatorio apresenta os resultados da verificacao de saude do Active Directory. ')
    p.add_run('A analise abrange conectividade de Domain Controllers, replicacao, DNS, ')
    p.add_run('politicas de grupo, contas de usuario, grupos privilegiados, seguranca e infraestrutura.')

    doc.add_paragraph()

    # Tabela de resumo executivo
    tabela_resumo = doc.add_table(rows=5, cols=2, style='Table Grid')
    tabela_resumo.alignment = WD_TABLE_ALIGNMENT.CENTER
    definir_bordas_tabela(tabela_resumo)

    dados_resumo = [
        ('Status Geral', saude_geral),
        ('Testes OK', str(total_ok)),
        ('Avisos', str(total_avisos)),
        ('Erros', str(total_erros)),
        ('Informativos', str(total_info)),
    ]

    cores_resumo = [
        None,
        CORES_BG['OK'],
        CORES_BG['AVISO'],
        CORES_BG['ERRO'],
        CORES_BG['INFO'],
    ]

    for i, (label, valor) in enumerate(dados_resumo):
        row = tabela_resumo.rows[i]
        row.cells[0].text = ''
        row.cells[1].text = ''

        p0 = row.cells[0].paragraphs[0]
        run0 = p0.add_run(label)
        run0.bold = True
        run0.font.size = Pt(10)

        p1 = row.cells[1].paragraphs[0]
        run1 = p1.add_run(valor)
        run1.bold = True
        run1.font.size = Pt(11)
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER

        if cores_resumo[i]:
            definir_cor_celula(row.cells[1], cores_resumo[i])

    # Largura das colunas
    for row in tabela_resumo.rows:
        row.cells[0].width = Cm(8)
        row.cells[1].width = Cm(8)

    # ========================================
    # 2. ESTATISTICAS / GRAFICOS
    # ========================================
    doc.add_paragraph()
    doc.add_heading('2. Estatisticas Gerais', level=1)

    # Gerar graficos
    dir_temp = Path(caminho_saida).parent
    img_pizza = str(dir_temp / '_temp_pizza.png')
    img_barras = str(dir_temp / '_temp_barras.png')

    graficos_gerados = []

    if gerar_grafico_pizza(contagem, img_pizza):
        doc.add_paragraph()
        doc.add_picture(img_pizza, width=Inches(4.5))
        last_paragraph = doc.paragraphs[-1]
        last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        graficos_gerados.append(img_pizza)

    if gerar_grafico_barras_categoria(dados, img_barras):
        doc.add_paragraph()
        doc.add_picture(img_barras, width=Inches(5.5))
        last_paragraph = doc.paragraphs[-1]
        last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        graficos_gerados.append(img_barras)

    if not graficos_gerados:
        adicionar_paragrafo(doc,
            "Graficos nao gerados (matplotlib nao disponivel).",
            tamanho=9, cor=RGBColor(0x99, 0x99, 0x99))

    # Tabela de contagem por categoria
    doc.add_paragraph()
    adicionar_paragrafo(doc, "Resumo por Categoria:", negrito=True, tamanho=11)

    cats_ordenadas = sorted(grupos.keys())
    tabela_cat = doc.add_table(rows=len(cats_ordenadas) + 1, cols=5, style='Table Grid')
    tabela_cat.alignment = WD_TABLE_ALIGNMENT.CENTER
    definir_bordas_tabela(tabela_cat)

    headers_cat = ['Categoria', 'OK', 'Avisos', 'Erros', 'Info']
    for i, h in enumerate(headers_cat):
        tabela_cat.rows[0].cells[i].text = h
    formatar_header_tabela(tabela_cat.rows[0], 5)

    for idx, cat in enumerate(cats_ordenadas):
        row = tabela_cat.rows[idx + 1]
        linhas_cat = grupos[cat]
        cont = Counter(r.get('Status', '').strip().upper() for r in linhas_cat)

        row.cells[0].text = cat
        row.cells[0].paragraphs[0].runs[0].font.size = Pt(9) if row.cells[0].paragraphs[0].runs else None

        for j, status in enumerate(['OK', 'AVISO', 'ERRO', 'INFO']):
            val = cont.get(status, 0)
            row.cells[j + 1].text = str(val)
            row.cells[j + 1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            if val > 0 and status in ('ERRO', 'AVISO'):
                definir_cor_celula(row.cells[j + 1], CORES_BG[status])

        # Linha alternada
        if idx % 2 == 1:
            for j in range(5):
                if not (row.cells[j]._tc.tcPr is not None and
                        row.cells[j]._tc.tcPr.findall(qn('w:shd'))):
                    definir_cor_celula(row.cells[j], COR_LINHA_ALTERNADA)

    doc.add_page_break()

    # ========================================
    # 3. ITENS CRITICOS (ERROS)
    # ========================================
    doc.add_heading('3. Itens Criticos (Erros)', level=1)

    erros = [r for r in dados if r.get('Status', '').strip().upper() == 'ERRO']

    if erros:
        adicionar_paragrafo(doc,
            f"Foram encontrados {len(erros)} item(ns) com status ERRO que requerem atencao imediata:",
            tamanho=10, espacamento_depois=6)

        tabela_erros = doc.add_table(rows=len(erros) + 1, cols=3, style='Table Grid')
        tabela_erros.alignment = WD_TABLE_ALIGNMENT.CENTER
        definir_bordas_tabela(tabela_erros)

        for i, h in enumerate(['Categoria', 'Teste', 'Detalhe']):
            tabela_erros.rows[0].cells[i].text = h
        formatar_header_tabela(tabela_erros.rows[0], 3)

        for idx, row_data in enumerate(erros):
            row = tabela_erros.rows[idx + 1]
            row.cells[0].text = row_data.get('Categoria', '')
            row.cells[1].text = row_data.get('Teste', '')
            row.cells[2].text = row_data.get('Detalhe', '')

            definir_cor_celula(row.cells[0], CORES_BG['ERRO'])

            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(9)
    else:
        adicionar_paragrafo(doc,
            "Nenhum item critico encontrado. Excelente!",
            tamanho=11, cor=CORES['OK'], negrito=True)

    doc.add_paragraph()

    # ========================================
    # 4. ITENS DE ATENCAO (AVISOS)
    # ========================================
    doc.add_heading('4. Itens de Atencao (Avisos)', level=1)

    avisos = [r for r in dados if r.get('Status', '').strip().upper() == 'AVISO']

    if avisos:
        adicionar_paragrafo(doc,
            f"Foram encontrados {len(avisos)} item(ns) com status AVISO para revisao:",
            tamanho=10, espacamento_depois=6)

        tabela_avisos = doc.add_table(rows=len(avisos) + 1, cols=3, style='Table Grid')
        tabela_avisos.alignment = WD_TABLE_ALIGNMENT.CENTER
        definir_bordas_tabela(tabela_avisos)

        for i, h in enumerate(['Categoria', 'Teste', 'Detalhe']):
            tabela_avisos.rows[0].cells[i].text = h
        formatar_header_tabela(tabela_avisos.rows[0], 3)

        for idx, row_data in enumerate(avisos):
            row = tabela_avisos.rows[idx + 1]
            row.cells[0].text = row_data.get('Categoria', '')
            row.cells[1].text = row_data.get('Teste', '')
            row.cells[2].text = row_data.get('Detalhe', '')

            definir_cor_celula(row.cells[0], CORES_BG['AVISO'])

            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(9)
    else:
        adicionar_paragrafo(doc,
            "Nenhum aviso encontrado.",
            tamanho=11, cor=CORES['OK'], negrito=True)

    doc.add_page_break()

    # ========================================
    # 5. RESULTADOS POR CATEGORIA
    # ========================================
    doc.add_heading('5. Resultados por Categoria', level=1)

    for cat in grupos:
        linhas_cat = grupos[cat]

        doc.add_heading(cat, level=2)

        # Mini resumo da categoria
        cont_cat = Counter(r.get('Status', '').strip().upper() for r in linhas_cat)
        resumo_parts = []
        for st in ['OK', 'AVISO', 'ERRO', 'INFO']:
            if cont_cat.get(st, 0) > 0:
                resumo_parts.append(f"{st}: {cont_cat[st]}")

        adicionar_paragrafo(doc, ' | '.join(resumo_parts),
                             tamanho=9, cor=RGBColor(0x66, 0x66, 0x66),
                             espacamento_depois=4)

        # Tabela da categoria
        tabela = doc.add_table(rows=len(linhas_cat) + 1, cols=3, style='Table Grid')
        tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
        definir_bordas_tabela(tabela)

        for i, h in enumerate(['Teste', 'Status', 'Detalhe']):
            tabela.rows[0].cells[i].text = h
        formatar_header_tabela(tabela.rows[0], 3)

        for idx, row_data in enumerate(linhas_cat):
            row = tabela.rows[idx + 1]
            status = row_data.get('Status', '').strip().upper()

            row.cells[0].text = row_data.get('Teste', '')
            row.cells[2].text = row_data.get('Detalhe', '')

            # Status com cor
            row.cells[1].text = ''
            p_status = row.cells[1].paragraphs[0]
            p_status.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_st = p_status.add_run(status)
            run_st.bold = True
            run_st.font.size = Pt(9)
            if status in CORES:
                run_st.font.color.rgb = CORES[status]

            # Fundo da linha para erros/avisos
            if status in CORES_BG and status in ('ERRO', 'AVISO'):
                for cell in row.cells:
                    definir_cor_celula(cell, CORES_BG[status])
            elif idx % 2 == 1:
                for cell in row.cells:
                    definir_cor_celula(cell, COR_LINHA_ALTERNADA)

            # Tamanho da fonte
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(9)

        # Ajustar largura
        for row in tabela.rows:
            row.cells[0].width = Cm(5)
            row.cells[1].width = Cm(2)
            row.cells[2].width = Cm(9)

    doc.add_page_break()

    # ========================================
    # 6. TABELA COMPLETA
    # ========================================
    doc.add_heading('6. Tabela Completa de Resultados', level=1)

    adicionar_paragrafo(doc,
        f"Total de {total} verificacoes realizadas.",
        tamanho=10, espacamento_depois=8)

    tabela_full = doc.add_table(rows=len(dados) + 1, cols=4, style='Table Grid')
    tabela_full.alignment = WD_TABLE_ALIGNMENT.CENTER
    definir_bordas_tabela(tabela_full)

    for i, h in enumerate(['Categoria', 'Teste', 'Status', 'Detalhe']):
        tabela_full.rows[0].cells[i].text = h
    formatar_header_tabela(tabela_full.rows[0], 4)

    cat_anterior = ""
    for idx, row_data in enumerate(dados):
        row = tabela_full.rows[idx + 1]
        cat = row_data.get('Categoria', '')
        status = row_data.get('Status', '').strip().upper()

        # Mostrar categoria apenas quando muda
        row.cells[0].text = cat if cat != cat_anterior else ''
        if cat != cat_anterior:
            for p in row.cells[0].paragraphs:
                for run in p.runs:
                    run.bold = True
        cat_anterior = cat

        row.cells[1].text = row_data.get('Teste', '')

        row.cells[2].text = ''
        p_st = row.cells[2].paragraphs[0]
        p_st.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_st = p_st.add_run(status)
        run_st.bold = True
        run_st.font.size = Pt(8)
        if status in CORES:
            run_st.font.color.rgb = CORES[status]

        row.cells[3].text = row_data.get('Detalhe', '')

        # Cor de fundo
        if status in ('ERRO', 'AVISO') and status in CORES_BG:
            for cell in row.cells:
                definir_cor_celula(cell, CORES_BG[status])
        elif idx % 2 == 1:
            for cell in row.cells:
                definir_cor_celula(cell, COR_LINHA_ALTERNADA)

        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    if not run.bold:
                        run.font.size = Pt(8)

    for row in tabela_full.rows:
        row.cells[0].width = Cm(3.5)
        row.cells[1].width = Cm(4)
        row.cells[2].width = Cm(1.5)
        row.cells[3].width = Cm(7)

    doc.add_page_break()

    # ========================================
    # 7. RECOMENDACOES
    # ========================================
    doc.add_heading('7. Recomendacoes', level=1)

    recomendacoes_map = {
        'AD Recycle Bin': 'Habilite o AD Recycle Bin para permitir restauracao facil de objetos excluidos: Enable-ADOptionalFeature "Recycle Bin Feature"',
        'LAPS': 'Implemente LAPS (Local Administrator Password Solution) para gerenciar senhas de admin local nas estacoes.',
        'Tombstone Lifetime': 'Aumente o Tombstone Lifetime para pelo menos 180 dias para evitar problemas de replicacao.',
        'Delegacao Irrestrita': 'Remova delegacao irrestrita de computadores que nao sao DCs. Use delegacao restrita (constrained) ou RBCD.',
        'Senha krbtgt': 'Resete a senha da conta krbtgt periodicamente (a cada 180 dias) para proteger tickets Kerberos.',
        'Conta Guest': 'Desabilite a conta Guest imediatamente.',
        'Nivel Funcional': 'Eleve o nivel funcional do dominio/floresta para aproveitar recursos de seguranca modernos.',
        'Kerberoast': 'Revise contas de usuario com SPNs. Use senhas longas (25+ chars) ou contas gerenciadas (gMSA).',
        'Senha nunca expira': 'Revise contas com senha que nunca expira. Remova essa configuracao quando possivel.',
        'Bloqueio': 'Configure politica de bloqueio de conta (lockout threshold) para proteger contra brute force.',
        'Complexidade': 'Habilite complexidade de senha e defina tamanho minimo de 12+ caracteres.',
        'GPO': 'Remova GPOs orfas (sem link) e desabilitadas que nao sao mais necessarias.',
        'Replicacao': 'Investigue e corrija falhas de replicacao imediatamente para evitar inconsistencias.',
        'NTP': 'Configure o PDC Emulator para sincronizar com uma fonte NTP externa confiavel.',
        'Disco': 'Monitore e expanda espaco em disco nos DCs que estao abaixo de 20% livre.',
        'OS Obsoleto': 'Planeje migracao de DCs e estacoes com sistemas operacionais fora de suporte.',
    }

    # Gerar recomendacoes baseadas nos erros e avisos encontrados
    recomendacoes_aplicaveis = []

    todos_textos = ' '.join(
        f"{r.get('Categoria','')} {r.get('Teste','')} {r.get('Detalhe','')}"
        for r in dados if r.get('Status', '').strip().upper() in ('ERRO', 'AVISO')
    ).upper()

    keywords_check = {
        'Recycle Bin':          ['RECYCLE BIN'],
        'LAPS':                 ['LAPS'],
        'Tombstone Lifetime':   ['TOMBSTONE'],
        'Delegacao Irrestrita': ['DELEGACAO', 'DELEGATION'],
        'Senha krbtgt':         ['KRBTGT'],
        'Conta Guest':          ['GUEST'],
        'Nivel Funcional':      ['NIVEL FUNCIONAL', '2008', '2003', '2012'],
        'Kerberoast':           ['KERBEROAST', 'SPN'],
        'Senha nunca expira':   ['NUNCA EXPIRA', 'NEVEREXPIRES'],
        'Bloqueio':             ['BLOQUEIO', 'LOCKOUT', 'THRESHOLD'],
        'Complexidade':         ['COMPLEXIDADE', 'COMPLEXITY'],
        'GPO':                  ['SEM LINK', 'ORFA', 'DESABILITADA'],
        'Replicacao':           ['REPLICACAO', 'REPLICATION', 'FALHA'],
        'NTP':                  ['FREE-RUNNING', 'LOCAL CMOS', 'NTP'],
        'Disco':                ['DISCO', 'DISK', '% LIVRE'],
        'OS Obsoleto':          ['OBSOLETO', 'DESATUALIZADO', '2008', '2003'],
    }

    for chave, keywords in keywords_check.items():
        if any(kw in todos_textos for kw in keywords):
            if chave in recomendacoes_map:
                recomendacoes_aplicaveis.append((chave, recomendacoes_map[chave]))

    if recomendacoes_aplicaveis:
        for i, (titulo, texto) in enumerate(recomendacoes_aplicaveis, 1):
            p = doc.add_paragraph()
            run_num = p.add_run(f"{i}. {titulo}: ")
            run_num.bold = True
            run_num.font.size = Pt(10)
            run_txt = p.add_run(texto)
            run_txt.font.size = Pt(10)
            p.paragraph_format.space_after = Pt(6)
    else:
        adicionar_paragrafo(doc,
            "Nenhuma recomendacao critica. O ambiente esta em boas condicoes.",
            tamanho=11, cor=CORES['OK'], negrito=True)

    # Recomendacoes gerais (sempre incluir)
    doc.add_paragraph()
    doc.add_heading('Boas Praticas Gerais', level=2)

    boas_praticas = [
        "Mantenha backups regulares do System State dos Domain Controllers.",
        "Monitore os Event Logs dos DCs (Directory Service, DNS Server, DFS Replication).",
        "Revise trimestralmente membros de grupos privilegiados (Domain Admins, Enterprise Admins).",
        "Implemente tiering model (Tier 0/1/2) para contas administrativas.",
        "Use PAWs (Privileged Access Workstations) para administracao do AD.",
        "Habilite auditoria avancada de logon e alteracoes no AD.",
        "Execute este health check mensalmente e compare com resultados anteriores.",
        "Documente todas as alteracoes realizadas no AD em um registro de mudancas.",
    ]

    for bp in boas_praticas:
        p = doc.add_paragraph(bp, style='List Bullet')
        for run in p.runs:
            run.font.size = Pt(10)

    # ========================================
    # RODAPE
    # ========================================
    doc.add_paragraph()
    doc.add_paragraph()

    p_footer = doc.add_paragraph()
    p_footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_f = p_footer.add_run(
        f"Relatorio gerado em {datetime.now().strftime('%d/%m/%Y %H:%mm:%S')} | "
        f"Health Check AD v5.0 | Documento Confidencial"
    )
    run_f.font.size = Pt(8)
    run_f.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

    # ========================================
    # SALVAR
    # ========================================
    doc.save(caminho_saida)

    # Limpar arquivos temporarios
    for img in graficos_gerados:
        try:
            os.remove(img)
        except OSError:
            pass

    return caminho_saida


# ============================================================
# MAIN
# ============================================================
def main():
    if len(sys.argv) < 2:
        print("=" * 50)
        print(" Gerador de Relatorio DOCX - Health Check AD")
        print("=" * 50)
        print()
        print("Uso: python gerar_relatorio_ad.py <arquivo.csv> [arquivo_saida.docx]")
        print()
        print("Exemplo:")
        print("  python gerar_relatorio_ad.py HealthCheckAD_20260416.csv")
        print("  python gerar_relatorio_ad.py HealthCheckAD_20260416.csv relatorio_ad.docx")
        print()
        print("O CSV deve ter as colunas: Categoria;Teste;Status;Detalhe")
        sys.exit(0)

    caminho_csv = sys.argv[1]

    if not os.path.isfile(caminho_csv):
        print(f"ERRO: Arquivo nao encontrado: {caminho_csv}")
        sys.exit(1)

    # Caminho de saida
    if len(sys.argv) >= 3:
        caminho_docx = sys.argv[2]
    else:
        caminho_docx = os.path.splitext(caminho_csv)[0] + '.docx'

    print("=" * 50)
    print(" Gerador de Relatorio DOCX - Health Check AD")
    print("=" * 50)
    print()
    print(f"  CSV de entrada:  {caminho_csv}")
    print(f"  DOCX de saida:   {caminho_docx}")
    print()

    # Ler dados
    print("[1/3] Lendo CSV...")
    dados = ler_csv(caminho_csv)
    contagem = contar_status(dados)
    print(f"       {len(dados)} registros encontrados")
    print(f"       OK={contagem.get('OK',0)} | AVISO={contagem.get('AVISO',0)} | ERRO={contagem.get('ERRO',0)} | INFO={contagem.get('INFO',0)}")
    print()

    # Gerar DOCX
    print("[2/3] Gerando documento DOCX...")
    resultado = gerar_docx(dados, caminho_docx)
    print()

    # Resultado
    print("[3/3] Concluido!")
    print()
    tamanho = os.path.getsize(caminho_docx)
    print(f"  Arquivo: {resultado}")
    print(f"  Tamanho: {tamanho / 1024:.1f} KB")
    print()
    print("=" * 50)


if __name__ == '__main__':
    main()